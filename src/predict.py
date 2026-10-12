import io
import sys
from functools import lru_cache
from pathlib import Path

import joblib
import numpy as np
from src import config
from src.preprocess import clean_text
from src.skills import extract_skills, fit_score

MAX_PDF_PAGES = 15
MIN_CHARS = 100
LOW_CONFIDENCE = 0.35


@lru_cache(maxsize=1)
def load_model(path=str(config.MODEL_PATH)):
    if not Path(path).exists():
        raise FileNotFoundError(f"Model not found at {path}. Run: python -m src.evaluate")
    return joblib.load(path)


def pretty_category(cat: str) -> str:
    return cat if len(cat) <= 3 else cat.replace("-", " ").title()


def version_warning():
    import json
    import sklearn
    info = Path(config.MODEL_PATH).with_name("model_info.json")
    try:
        trained = json.loads(info.read_text(encoding="utf-8"))["sklearn"]
    except (OSError, KeyError, ValueError):
        return None
    if trained != sklearn.__version__:
        return (f"Model was trained with scikit-learn {trained} but {sklearn.__version__} is installed. "
                f"Predictions may be unreliable. Install scikit-learn=={trained} (see requirements.txt) "
                "or retrain with `python -m src.evaluate`.")
    return None


def extract_text(source, filename=None) -> str:
    name = (filename or getattr(source, "name", None) or str(source)).lower()
    try:
        if isinstance(source, (str, Path)):
            data = Path(source).read_bytes()
        else:
            if hasattr(source, "seek"):
                source.seek(0)
            data = source.read()
        if name.endswith(".pdf"):
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                return "\n".join((page.extract_text() or "") for page in pdf.pages[:MAX_PDF_PAGES])
        if name.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                parts += [c.text for row in table.rows for c in row.cells]
            return "\n".join(parts)
        if name.endswith(".txt"):
            try:
                return data.decode("utf-8")
            except UnicodeDecodeError:
                return data.decode("cp1252", errors="replace")
    except Exception as e:
        raise ValueError(f"Could not read '{name}': {e}") from e
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def predict_resume(text, jd_text=None, pipe=None, top_k=3) -> dict:
    pipe = pipe or load_model()
    text = text if isinstance(text, str) else ""
    warnings = []

    if len(text.strip()) < MIN_CHARS:
        warnings.append("Very little text was extracted (scanned/image PDF or empty file?). "
                        "OCR is not supported, so predictions are unreliable.")
        return {"top_predictions": [], "skills": extract_skills(text), "fit": None,
                "warnings": warnings, "n_chars": len(text.strip())}

    letters = [c for c in text if c.isalpha()]
    non_ascii = sum(ord(c) > 127 for c in letters) / len(letters) if letters else 0.0
    if non_ascii > 0.3:
        warnings.append("This text is mostly non-Latin script. The model was trained on English "
                        "resumes only, so predictions are not meaningful.")

    cleaned = clean_text(text)
    if not cleaned:
        if not non_ascii > 0.3:
            warnings.append("No usable words remained after cleaning the text.")
        return {"top_predictions": [], "skills": [], "fit": None,
                "warnings": warnings, "n_chars": len(text.strip())}

    proba = pipe.predict_proba([cleaned])[0]
    order = np.argsort(proba)[::-1][:top_k]
    top = [{"category": str(pipe.classes_[i]), "probability": float(proba[i])} for i in order]
    if top[0]["probability"] < LOW_CONFIDENCE:
        warnings.append(f"Low confidence (top probability {top[0]['probability']:.2f}). "
                        "Consider the top-3 categories rather than only the first.")
    if len(text.split()) < 150:
        warnings.append("This resume is short; predictions may be less reliable.")

    result = {"top_predictions": top, "skills": extract_skills(text),
              "fit": None, "warnings": warnings, "n_chars": len(text.strip())}
    if jd_text and jd_text.strip():
        result["fit"] = fit_score(text, jd_text, pipe)
    return result


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("usage: python -m src.predict resume.(pdf|docx|txt) [jd.txt]")
    resume = extract_text(sys.argv[1])
    jd = extract_text(sys.argv[2]) if len(sys.argv) > 2 else None
    out = predict_resume(resume, jd)
    for t in out["top_predictions"]:
        print(f"{t['category']:25s} {t['probability']:.2f}")
    print("skills:", ", ".join(out["skills"]))
    if out["fit"]:
        print(f"fit: {out['fit']['fit']:.3f}  matched: {out['fit']['matched_skills']}  missing: {out['fit']['missing_skills']}")
    for w in out["warnings"]:
        print("WARNING:", w)
